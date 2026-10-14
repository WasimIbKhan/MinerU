"""Local patch (local-patches-v4): a reused DOCX list only becomes headings when its
numbering levels carry an outline level (mineru/parser/docx_lists.py)."""
from collections import Counter

import pytest

docx = pytest.importorskip("docx")
from docx.oxml import parse_xml  # noqa: E402

from mineru.parser import parse  # noqa: E402

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _shared_list(doc, fmt: str, styles: tuple[str | None, str | None]) -> int:
    numbering = doc.part.numbering_part.element
    levels = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
        + (f'<w:pStyle w:val="{style}"/>' if style else "")
        + f'<w:lvlText w:val="{"%" + str(i + 1) + "." if fmt == "decimal" else "o"}"/></w:lvl>'
        for i, style in enumerate(styles)
    )
    abstract_id = 90
    numbering.insert(0, parse_xml(f'<w:abstractNum xmlns:w="{_W}" w:abstractNumId="{abstract_id}">'
                                  f'<w:multiLevelType w:val="multilevel"/>{levels}</w:abstractNum>'))
    return numbering.add_num(abstract_id).numId


def _item(doc, text: str, num_id: int, ilvl: int, style: str = "List Paragraph"):
    paragraph = doc.add_paragraph(text, style=style)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = ilvl
    num_pr.get_or_add_numId().val = num_id


def _block_counts(tmp_path, build) -> Counter:
    doc = docx.Document()
    build(doc)
    path = tmp_path / "doc.docx"
    doc.save(path)
    content = parse(str(path), tier="flash").structured_content()
    return Counter((b["type"], b.get("level")) for p in content["pages"] for b in p["blocks"]
                   if b["type"] in ("paragraph_title", "list"))


def test_bullet_list_reused_between_body_paragraphs_stays_a_list(tmp_path):
    def build(doc):
        num_id = _shared_list(doc, "bullet", (None, None))
        doc.add_heading("Part One", level=1)
        doc.add_paragraph("Opening paragraph of the part.")
        _item(doc, "First point made as a full sentence.", num_id, 0)
        _item(doc, "A nested remark about the first point.", num_id, 1)
        doc.add_paragraph("A body paragraph between the two bullet groups.")
        doc.add_heading("A section", level=2)
        _item(doc, "Second point made as a full sentence.", num_id, 0)
        _item(doc, "Another nested remark.", num_id, 1)
        doc.add_paragraph("Closing paragraph.")

    counts = _block_counts(tmp_path, build)
    assert counts[("paragraph_title", 2)] == 1 and counts[("paragraph_title", 3)] == 1
    assert sum(n for (kind, _), n in counts.items() if kind == "paragraph_title") == 2
    assert counts[("list", None)] == 2


def test_numbered_list_linked_to_heading_styles_keeps_its_headings(tmp_path):
    def build(doc):
        num_id = _shared_list(doc, "decimal", ("Heading1", "Heading2"))
        _item(doc, "Introduction", num_id, 0, style="Normal")
        doc.add_paragraph("Body text of the introduction.")
        _item(doc, "Background", num_id, 1, style="Normal")
        doc.add_paragraph("Body text of the background.")
        _item(doc, "Method", num_id, 0, style="Normal")
        doc.add_paragraph("Body text of the method.")

    counts = _block_counts(tmp_path, build)
    assert counts[("paragraph_title", 2)] == 2 and counts[("paragraph_title", 3)] == 1
    assert counts[("list", None)] == 0
