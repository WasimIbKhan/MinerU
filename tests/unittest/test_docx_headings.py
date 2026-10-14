"""Local patch (local-patches-v4): DOCX headings from inherited styles, and
headings that end an open list (mineru/parser/docx_headings.py)."""
import pytest

docx = pytest.importorskip("docx")
from docx.oxml import parse_xml  # noqa: E402

from _docx_test_utils import W_NS, parse_blocks  # noqa: E402


def test_heading_based_on_a_localised_heading_style_is_a_heading(tmp_path):
    doc = docx.Document()
    styles = doc.styles.element
    styles.append(parse_xml(
        f'<w:style xmlns:w="{W_NS}" w:type="paragraph" w:styleId="2"><w:name w:val="heading 2"/>'
        '<w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style>'))
    styles.append(parse_xml(
        f'<w:style xmlns:w="{W_NS}" w:type="paragraph" w:customStyle="1" w:styleId="Myheading2">'
        '<w:name w:val="Myheading2"/><w:basedOn w:val="2"/></w:style>'))
    doc.add_paragraph("Chapter heading", style="Myheading2")
    doc.add_paragraph("Body text.")
    path = tmp_path / "inherited_heading.docx"
    doc.save(path)

    assert [(b["type"], b.get("level")) for b in parse_blocks(path)] == [("paragraph_title", 3), ("text", None)]


def test_list_items_after_a_heading_stay_after_it(tmp_path):
    doc = docx.Document()
    doc.add_heading("First section", level=1)
    doc.add_paragraph("Opening item", style="List Bullet")
    doc.add_heading("Second section", level=1)
    doc.add_paragraph("Item under the second heading", style="List Bullet")
    path = tmp_path / "list_after_heading.docx"
    doc.save(path)

    blocks = parse_blocks(path)
    assert [b["type"] for b in blocks] == ["paragraph_title", "list", "paragraph_title", "list"]
    assert "Item under the second heading" in blocks[3]["content"]
