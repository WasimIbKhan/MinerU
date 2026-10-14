"""Local patch (local-patches-v4): DOCX package repairs before docvortex reads it
(mineru/parser/docx_package.py)."""
import base64

import pytest

docx = pytest.importorskip("docx")
from docx.oxml import parse_xml  # noqa: E402

from _docx_test_utils import W_NS, parse_blocks, write_package  # noqa: E402

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _wrap_in_content_control(elements) -> None:
    sdt = parse_xml(f'<w:sdt xmlns:w="{W_NS}"><w:sdtPr><w:tag w:val="goog_rdk_1"/></w:sdtPr><w:sdtContent/></w:sdt>')
    elements[0].addprevious(sdt)
    content = sdt.find(f"{{{W_NS}}}sdtContent")
    for element in elements:
        content.append(element)


def test_wordprocessingml_under_another_prefix_parses(tmp_path):
    document = (
        f'<ns0:document xmlns:ns0="{W_NS}"><ns0:body>'
        '<ns0:p><ns0:pPr><ns0:numPr><ns0:ilvl ns0:val="0"/><ns0:numId ns0:val="0"/></ns0:numPr></ns0:pPr>'
        "<ns0:r><ns0:t>Written by ElementTree</ns0:t></ns0:r></ns0:p>"
        "</ns0:body></ns0:document>"
    )
    path = tmp_path / "ns0.docx"
    write_package(path, document)

    assert [(b["type"], b["content"]) for b in parse_blocks(path)] == [("text", "Written by ElementTree")]


def test_table_inside_a_content_control_stays_a_table(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Before the table.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"cell {r}{c}"
    doc.add_paragraph("After the table.")
    _wrap_in_content_control([table._tbl])
    path = tmp_path / "sdt_table.docx"
    doc.save(path)

    blocks = parse_blocks(path)
    assert [b["type"] for b in blocks] == ["text", "table", "text"]
    assert all(f"cell {r}{c}" in blocks[1]["content"] for r in range(2) for c in range(2))


def test_text_beside_a_picture_inside_a_content_control_is_kept(tmp_path):
    image = tmp_path / "pixel.png"
    image.write_bytes(_PNG)
    doc = docx.Document()
    picture = doc.add_paragraph()
    picture.add_run().add_picture(str(image))
    title = doc.add_paragraph("Cover title inside the control")
    doc.add_paragraph("Body text outside the control")
    _wrap_in_content_control([picture._p, title._p])
    path = tmp_path / "sdt_picture.docx"
    doc.save(path)

    texts = [b["content"] for b in parse_blocks(path) if b["type"] == "text"]
    assert texts == ["Cover title inside the control", "Body text outside the control"]
